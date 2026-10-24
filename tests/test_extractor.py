import pytest
import pymupdf
from docx import Document
from organizer.extractor import extract_text, ExtractionError


def test_txt_bom_and_truncation(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("\ufeffUniversity lecture notes", encoding="utf-8")
    assert extract_text(path, 10).text == "University"


def test_invalid_utf8(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_bytes(b"hello\xff")
    assert extract_text(path).warning


def test_pdf(tmp_path):
    path = tmp_path / "invoice.pdf"
    with pymupdf.open() as doc:
        doc.new_page().insert_text((72, 72), "Invoice amount: 100")
        doc.save(path)
    assert "Invoice amount" in extract_text(path).text


def test_blank_pdf(tmp_path):
    path = tmp_path / "scan.pdf"
    with pymupdf.open() as doc:
        doc.new_page()
        doc.save(path)
    assert extract_text(path).text == ""
    assert "no OCR" in extract_text(path).warning


def test_docx_paragraph_and_table(tmp_path):
    path = tmp_path / "contract.docx"
    doc = Document()
    doc.add_paragraph("Employment agreement")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Terms"
    doc.save(path)
    assert "Terms" in extract_text(path).text
    assert "Employment agreement" in extract_text(path).text


@pytest.mark.parametrize("suffix", [".pdf", ".docx"])
def test_corrupt(tmp_path, suffix):
    path = tmp_path / ("broken" + suffix)
    path.write_bytes(b"invalid")
    with pytest.raises(ExtractionError):
        extract_text(path)


def test_symlink(tmp_path):
    target = tmp_path / "target.txt"
    target.touch()
    link = tmp_path / "link.txt"
    link.symlink_to(target)
    with pytest.raises(ValueError):
        extract_text(link)
